"""Bounded document extraction for supported learning materials."""
from __future__ import annotations
import io, os, re, tempfile, subprocess, ipaddress, socket
from urllib.parse import urlparse, parse_qs
import requests
from bs4 import BeautifulSoup
from config import MAX_EXTRACTED_CHARS, MAX_MEDIA_SECONDS


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".pptx", ".mp3", ".wav", ".m4a", ".mp4", ".mov", ".webm", ".mpeg", ".mpga"}

class IngestionError(ValueError):
    pass

def _limit(text: str) -> str:
    text = re.sub(r"\x00", "", text or "").strip()
    if not text:
        raise IngestionError("No readable text was found in this source.")
    return text[:MAX_EXTRACTED_CHARS]

def _pdf(data: bytes) -> str:
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        return "\n\n".join(f"[Page {i+1}]\n{page.get_text()}" for i, page in enumerate(doc))
    except Exception as exc:
        raise IngestionError("Could not read this PDF. It may be damaged or image-only; OCR is not enabled in v1.") from exc

def _docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(parts)
    except Exception as exc:
        raise IngestionError("Could not parse this Word document.") from exc

def _pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        return "\n\n".join(f"[Slide {i+1}]\n" + "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()) for i, slide in enumerate(prs.slides))
    except Exception as exc:
        raise IngestionError("Could not parse this PowerPoint file.") from exc

def _media(data: bytes, filename: str, ext: str) -> str:
    # Validate duration with ffprobe before invoking Whisper.
    with tempfile.TemporaryDirectory(prefix="rag-media-") as tmp:
        source = os.path.join(tmp, "input" + ext)
        with open(source, "wb") as f: f.write(data)
        try:
            probe = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=noprint_wrappers=1:nokey=1", source], capture_output=True, text=True, timeout=20, check=True)
            duration = float(probe.stdout.strip())
        except Exception as exc:
            raise IngestionError("Media inspection failed. Ensure FFmpeg and ffprobe are installed and the file is valid.") from exc
        if duration > MAX_MEDIA_SECONDS:
            raise IngestionError(f"Media exceeds the {MAX_MEDIA_SECONDS // 60}-minute duration limit.")
        try:
            from groq import Groq
            from config import GROQ_API_KEY, WHISPER_MODEL
            if not GROQ_API_KEY:
                raise IngestionError("GROQ_API_KEY is missing. Add it to Backend/.env before transcribing audio or video.")

            # Groq's speech-to-text endpoint supports a defined list of audio/video formats.
            # Extract a compact MP3 from video containers (including MOV) for consistent support.
            upload_path = source
            if ext in {".mp4", ".mov", ".webm", ".mpeg"}:
                upload_path = os.path.join(tmp, "transcription-audio.mp3")
                try:
                    subprocess.run(
                        ["ffmpeg", "-y", "-v", "error", "-i", source, "-vn", "-ac", "1", "-ar", "16000", "-b:a", "64k", upload_path],
                        capture_output=True, text=True, timeout=90, check=True,
                    )
                except Exception as exc:
                    raise IngestionError("Could not extract audio from this video. Check that it contains an audio track and FFmpeg is installed.") from exc
            if os.path.getsize(upload_path) > 25 * 1024 * 1024:
                raise IngestionError("The extracted audio exceeds Groq's 25 MB transcription upload limit. Use a shorter or more compressed file.")

            client = Groq(api_key=GROQ_API_KEY, timeout=90.0, max_retries=2)
            with open(upload_path, "rb") as media_file:
                result = client.audio.transcriptions.create(
                    file=(os.path.basename(upload_path), media_file.read()),
                    model=WHISPER_MODEL,
                    response_format="verbose_json",
                    timestamp_granularities=["segment"],
                    temperature=0.0,
                )
            segments = getattr(result, "segments", None) or []
            lines = []
            for seg in segments:
                if isinstance(seg, dict):
                    start, end, segment_text = seg.get("start", 0), seg.get("end", 0), seg.get("text", "")
                else:
                    start, end, segment_text = getattr(seg, "start", 0), getattr(seg, "end", 0), getattr(seg, "text", "")
                start, end = max(0, int(float(start))), max(0, int(float(end)))
                if str(segment_text).strip():
                    lines.append(f"[Timestamp {start//60:02d}:{start%60:02d} - {end//60:02d}:{end%60:02d}] {str(segment_text).strip()}")
            transcript = "\n".join(lines).strip() or getattr(result, "text", "") or ""
            if not transcript.strip():
                raise IngestionError("Groq Whisper returned an empty transcript. Try a clearer recording or another supported file format.")
            return transcript
        except IngestionError:
            raise
        except Exception as exc:
            # Do not leak provider response details or credentials to the UI.
            raise IngestionError("Groq Whisper transcription failed. Check GROQ_API_KEY, model availability, API limits, and file format.") from exc

def extract_file(data: bytes, filename: str) -> tuple[str, str]:
    ext = os.path.splitext(filename.lower())[1]
    if ext not in ALLOWED_EXTENSIONS:
        raise IngestionError("Unsupported file type. Supported: PDF, DOCX, TXT, MD, PPTX, audio and video.")
    if ext == ".pdf": text = _pdf(data); source_type = "pdf"
    elif ext == ".docx": text = _docx(data); source_type = "docx"
    elif ext == ".pptx": text = _pptx(data); source_type = "pptx"
    elif ext in {".txt", ".md"}:
        try: text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc: raise IngestionError("Text file must use UTF-8 encoding.") from exc
        source_type = "text"
    else:
        text = _media(data, filename, ext); source_type = "video" if ext in {".mp4", ".mov", ".webm", ".mpeg"} else "audio"
    return _limit(text), source_type

def _public_ip(host: str) -> bool:
    try:
        infos = socket.getaddrinfo(host, None, type=socket.SOCK_STREAM)
        return bool(infos) and all(ipaddress.ip_address(info[4][0]).is_global for info in infos)
    except Exception:
        return False


def _youtube_video_id(parsed):
    """Extract a video ID from common YouTube URL formats."""
    host = (parsed.hostname or "").lower()
    path_parts = [part for part in parsed.path.split("/") if part]

    video_id = None

    if host in {"youtu.be", "www.youtu.be"}:
        video_id = path_parts[0] if path_parts else None
    elif host in {
        "youtube.com",
        "www.youtube.com",
        "m.youtube.com",
        "youtube-nocookie.com",
        "www.youtube-nocookie.com",
    }:
        if parsed.path == "/watch":
            video_id = parse_qs(parsed.query).get("v", [None])[0]
        elif len(path_parts) >= 2 and path_parts[0] in {
            "shorts", "embed", "live"
        }:
            video_id = path_parts[1]

    if not video_id or not re.fullmatch(r"[A-Za-z0-9_-]{11}", video_id):
        raise IngestionError(
            "Could not find a valid YouTube video ID. "
            "Submit a direct video URL, not a playlist URL."
        )

    return video_id


def _extract_youtube(url, parsed):
    """Extract available YouTube captions and preserve video timestamps."""
    video_id = _youtube_video_id(parsed)

    try:
        from youtube_transcript_api import YouTubeTranscriptApi

        api = YouTubeTranscriptApi()
        transcript_list = api.list(video_id)

        # Prefer English, then Hindi, then another available caption language.
        try:
            transcript = transcript_list.find_transcript(["en"])
        except Exception:
            try:
                transcript = transcript_list.find_transcript(["hi"])
            except Exception:
                transcript = next(iter(transcript_list), None)

        if transcript is None:
            raise IngestionError(
                "No captions or transcripts are available for this YouTube video."
            )

        snippets = transcript.fetch()
        lines = []

        for item in snippets:
            start = float(item.start)
            end = start + float(item.duration)
            text = str(item.text).replace("\n", " ").strip()

            if not text:
                continue

            start_seconds = max(0, int(start))
            end_seconds = max(start_seconds, int(end))

            start_time = (
                f"{start_seconds // 60:02d}:{start_seconds % 60:02d}"
            )
            end_time = (
                f"{end_seconds // 60:02d}:{end_seconds % 60:02d}"
            )

            lines.append(
                f"[Timestamp {start_time} - {end_time}] {text}"
            )

        if not lines:
            raise IngestionError(
                "The YouTube transcript is empty."
            )

        # Use YouTube's official oEmbed endpoint for the title.
        title = f"YouTube video {video_id}"
        try:
            response = requests.get(
                "https://www.youtube.com/oembed",
                params={
                    "url": f"https://www.youtube.com/watch?v={video_id}",
                    "format": "json",
                },
                timeout=(3, 5),
                allow_redirects=False,
            )
            if response.ok:
                title = response.json().get("title") or title
        except (requests.RequestException, ValueError):
            pass

        final_url = f"https://www.youtube.com/watch?v={video_id}"
        return _limit("\n".join(lines)), title[:300], final_url

    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(
            "Could not retrieve this YouTube transcript. "
            "Check whether captions are available and accessible. "
            "Some videos restrict transcript access."
        ) from exc


def extract_url(url: str) -> tuple[str, str, str]:
    parsed = urlparse(url)

    if (
        parsed.scheme not in {"http", "https"}
        or not parsed.hostname
        or parsed.username
        or parsed.password
    ):
        raise IngestionError(
            "Enter a valid public HTTP or HTTPS URL."
        )

    host = (parsed.hostname or "").lower()

    youtube_hosts = {
        "youtube.com",
        "www.youtube.com",
        "m.youtube.com",
        "youtu.be",
        "www.youtu.be",
        "youtube-nocookie.com",
        "www.youtube-nocookie.com",
    }

    if host in youtube_hosts:
        if parsed.port and parsed.port not in {80, 443}:
            raise IngestionError(
                "Only standard HTTP/HTTPS ports are allowed."
            )
        return _extract_youtube(url, parsed)

    # Existing public-webpage ingestion remains supported.
    if parsed.port and parsed.port not in {80, 443}:
        raise IngestionError(
            "Only standard HTTP/HTTPS ports are allowed."
        )

    if not _public_ip(parsed.hostname):
        raise IngestionError(
            "This host is not publicly accessible or could not be safely resolved."
        )

    try:
        resp = requests.get(
            url,
            timeout=(5, 15),
            allow_redirects=False,
            headers={"User-Agent": "StudyAssistant/1.0"},
            stream=True,
        )

        if 300 <= resp.status_code < 400:
            raise IngestionError(
                "Redirects are not followed for security. "
                "Submit the destination URL directly."
            )

        resp.raise_for_status()

        if "text/html" not in resp.headers.get("Content-Type", "").lower():
            raise IngestionError(
                "Only public HTML pages are supported for URL ingestion."
            )

        chunks = []
        total = 0

        for chunk in resp.iter_content(16384):
            total += len(chunk)

            if total > 3 * 1024 * 1024:
                raise IngestionError(
                    "Web page exceeds the 3 MB download limit."
                )

            chunks.append(chunk)

        html = b"".join(chunks).decode(
            resp.encoding or "utf-8",
            errors="replace",
        )

        soup = BeautifulSoup(html, "html.parser")

        for element in soup(
            ["script", "style", "nav", "footer", "header",
             "noscript", "form"]
        ):
            element.decompose()

        title = (
            soup.title.get_text(" ", strip=True)
            if soup.title
            else parsed.hostname
        )

        main = soup.find("main") or soup.find("article") or soup.body or soup

        return (
            _limit(main.get_text("\n", strip=True)),
            title[:300],
            url,
        )

    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(
            "Could not fetch this page. It may block automated access "
            "or be temporarily unavailable."
        ) from exc

